"""
文档解析服务
支持多种文档格式的解析和转换
"""

import os
import logging
from typing import Dict, Any, Optional
from pathlib import Path

logger = logging.getLogger(__name__)


class DocumentParser:
    """文档解析器基类"""
    
    @staticmethod
    def parse(file_path: str) -> Dict[str, Any]:
        """
        解析文档文件
        
        Args:
            file_path: 文件路径
            
        Returns:
            {
                'success': bool,
                'content': str,  # 文本内容
                'html': str,     # HTML格式（如果支持）
                'metadata': dict,  # 元数据
                'error': str     # 错误信息
            }
        """
        raise NotImplementedError


class WordParser(DocumentParser):
    """Word文档解析器 (.docx)"""
    
    @staticmethod
    def parse(file_path: str) -> Dict[str, Any]:
        try:
            from docx import Document
            from docx.oxml.text.paragraph import CT_P
            from docx.oxml.table import CT_Tbl
            from docx.table import _Cell, Table
            from docx.text.paragraph import Paragraph
            
            doc = Document(file_path)
            
            # 提取文本内容
            text_content = []
            html_content = []
            
            for element in doc.element.body:
                if isinstance(element, CT_P):
                    # 段落
                    para = Paragraph(element, doc)
                    text = para.text.strip()
                    if text:
                        text_content.append(text)
                        
                        # 转换为HTML
                        if para.style.name.startswith('Heading'):
                            level = para.style.name.replace('Heading ', '')
                            html_content.append(f'<h{level}>{text}</h{level}>')
                        else:
                            html_content.append(f'<p>{text}</p>')
                
                elif isinstance(element, CT_Tbl):
                    # 表格
                    table = Table(element, doc)
                    table_text = []
                    table_html = ['<table>']
                    
                    for row in table.rows:
                        row_text = []
                        table_html.append('<tr>')
                        for cell in row.cells:
                            cell_text = cell.text.strip()
                            row_text.append(cell_text)
                            table_html.append(f'<td>{cell_text}</td>')
                        table_html.append('</tr>')
                        table_text.append(' | '.join(row_text))
                    
                    table_html.append('</table>')
                    text_content.append('\n'.join(table_text))
                    html_content.append(''.join(table_html))
            
            # 提取元数据
            metadata = {
                'title': doc.core_properties.title or '',
                'author': doc.core_properties.author or '',
                'created': str(doc.core_properties.created) if doc.core_properties.created else '',
                'modified': str(doc.core_properties.modified) if doc.core_properties.modified else '',
                'paragraphs': len(doc.paragraphs),
                'tables': len(doc.tables)
            }
            
            return {
                'success': True,
                'content': '\n\n'.join(text_content),
                'html': '\n'.join(html_content),
                'metadata': metadata,
                'error': None
            }
            
        except ImportError:
            return {
                'success': False,
                'content': '',
                'html': '',
                'metadata': {},
                'error': '缺少python-docx库，请安装: pip install python-docx'
            }
        except Exception as e:
            logger.error(f"解析Word文档失败: {e}", exc_info=True)
            return {
                'success': False,
                'content': '',
                'html': '',
                'metadata': {},
                'error': f'解析失败: {str(e)}'
            }


class PDFParser(DocumentParser):
    """PDF文档解析器"""
    
    @staticmethod
    def parse(file_path: str) -> Dict[str, Any]:
        try:
            import PyPDF2
            
            text_content = []
            metadata = {}
            
            with open(file_path, 'rb') as file:
                pdf_reader = PyPDF2.PdfReader(file)
                
                # 提取元数据
                if pdf_reader.metadata:
                    metadata = {
                        'title': pdf_reader.metadata.get('/Title', ''),
                        'author': pdf_reader.metadata.get('/Author', ''),
                        'subject': pdf_reader.metadata.get('/Subject', ''),
                        'creator': pdf_reader.metadata.get('/Creator', ''),
                        'pages': len(pdf_reader.pages)
                    }
                
                # 提取文本
                for page_num, page in enumerate(pdf_reader.pages):
                    try:
                        text = page.extract_text()
                        if text.strip():
                            text_content.append(f"--- 第 {page_num + 1} 页 ---\n{text}")
                    except Exception as e:
                        logger.warning(f"提取第{page_num + 1}页失败: {e}")
            
            content = '\n\n'.join(text_content)
            
            # 简单的HTML转换
            html_content = content.replace('\n', '<br>')
            
            return {
                'success': True,
                'content': content,
                'html': f'<div>{html_content}</div>',
                'metadata': metadata,
                'error': None
            }
            
        except ImportError:
            return {
                'success': False,
                'content': '',
                'html': '',
                'metadata': {},
                'error': '缺少PyPDF2库，请安装: pip install PyPDF2'
            }
        except Exception as e:
            logger.error(f"解析PDF文档失败: {e}", exc_info=True)
            return {
                'success': False,
                'content': '',
                'html': '',
                'metadata': {},
                'error': f'解析失败: {str(e)}'
            }


class MarkdownParser(DocumentParser):
    """Markdown文档解析器"""
    
    @staticmethod
    def parse(file_path: str) -> Dict[str, Any]:
        try:
            with open(file_path, 'r', encoding='utf-8') as file:
                content = file.read()
            
            # 尝试转换为HTML
            try:
                import markdown
                html_content = markdown.markdown(content, extensions=['extra', 'codehilite'])
            except ImportError:
                # 如果没有markdown库，使用简单转换
                html_content = content.replace('\n', '<br>')
            
            # 提取元数据（如果有YAML front matter）
            metadata = {}
            if content.startswith('---'):
                try:
                    import yaml
                    parts = content.split('---', 2)
                    if len(parts) >= 3:
                        metadata = yaml.safe_load(parts[1])
                        content = parts[2].strip()
                except:
                    pass
            
            return {
                'success': True,
                'content': content,
                'html': html_content,
                'metadata': metadata,
                'error': None
            }
            
        except Exception as e:
            logger.error(f"解析Markdown文档失败: {e}", exc_info=True)
            return {
                'success': False,
                'content': '',
                'html': '',
                'metadata': {},
                'error': f'解析失败: {str(e)}'
            }


class TextParser(DocumentParser):
    """纯文本文档解析器"""
    
    @staticmethod
    def parse(file_path: str) -> Dict[str, Any]:
        try:
            # 尝试多种编码
            encodings = ['utf-8', 'gbk', 'gb2312', 'latin-1']
            content = None
            
            for encoding in encodings:
                try:
                    with open(file_path, 'r', encoding=encoding) as file:
                        content = file.read()
                    break
                except UnicodeDecodeError:
                    continue
            
            if content is None:
                raise ValueError('无法识别文件编码')
            
            # 简单的HTML转换
            html_content = content.replace('\n', '<br>')
            
            # 获取文件信息
            file_stat = os.stat(file_path)
            metadata = {
                'size': file_stat.st_size,
                'created': str(file_stat.st_ctime),
                'modified': str(file_stat.st_mtime)
            }
            
            return {
                'success': True,
                'content': content,
                'html': f'<div>{html_content}</div>',
                'metadata': metadata,
                'error': None
            }
            
        except Exception as e:
            logger.error(f"解析文本文档失败: {e}", exc_info=True)
            return {
                'success': False,
                'content': '',
                'html': '',
                'metadata': {},
                'error': f'解析失败: {str(e)}'
            }


class DocumentParserService:
    """文档解析服务"""
    
    # 支持的文件格式
    PARSERS = {
        '.docx': WordParser,
        '.doc': WordParser,  # 需要额外处理
        '.pdf': PDFParser,
        '.md': MarkdownParser,
        '.markdown': MarkdownParser,
        '.txt': TextParser,
        '.text': TextParser,
    }
    
    @classmethod
    def parse_document(cls, file_path: str) -> Dict[str, Any]:
        """
        解析文档
        
        Args:
            file_path: 文件路径
            
        Returns:
            解析结果
        """
        if not os.path.exists(file_path):
            return {
                'success': False,
                'content': '',
                'html': '',
                'metadata': {},
                'error': '文件不存在'
            }
        
        # 获取文件扩展名
        ext = Path(file_path).suffix.lower()
        
        # 查找对应的解析器
        parser_class = cls.PARSERS.get(ext)
        
        if not parser_class:
            return {
                'success': False,
                'content': '',
                'html': '',
                'metadata': {},
                'error': f'不支持的文件格式: {ext}'
            }
        
        # 解析文档
        return parser_class.parse(file_path)
    
    @classmethod
    def get_supported_formats(cls) -> list:
        """获取支持的文件格式列表"""
        return list(cls.PARSERS.keys())
    
    @classmethod
    def is_supported(cls, file_path: str) -> bool:
        """检查文件格式是否支持"""
        ext = Path(file_path).suffix.lower()
        return ext in cls.PARSERS


# 全局服务实例
document_parser_service = DocumentParserService()
